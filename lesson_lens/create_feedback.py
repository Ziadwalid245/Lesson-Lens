"""Turns a StudentFeedback object into a tidy Word document, or plain text for email.

Each lesson folder keeps feedback.json next to feedback.docx, so the teacher
can edit the feedback in the app later and export it again.
"""
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .feedback_structure import StudentFeedback


def _table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    doc.add_paragraph()  # breathing room after the table


def _bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _date_line(lesson_date):
    return (lesson_date or date.today()).strftime("%A %d %B %Y")


def create_feedback_doc(feedback, output_path, lesson_date=None):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Lesson Feedback", level=0)
    doc.add_paragraph(_date_line(lesson_date))

    doc.add_heading("What we covered", level=1)
    doc.add_paragraph(feedback.lesson_summary)

    if feedback.positive_feedback:
        doc.add_heading("What you did well", level=1)
        _bullets(doc, feedback.positive_feedback)

    if feedback.grammar_points:
        doc.add_heading("Grammar", level=1)
        _table(doc, ["Grammar point", "Form", "Usage"],
               [(g.name, g.form, g.usage) for g in feedback.grammar_points])

    if feedback.vocab_items:
        doc.add_heading("Vocabulary", level=1)
        _table(doc, ["Word", "Meaning", "From the lesson"],
               [(v.word, v.definition, v.in_context) for v in feedback.vocab_items])

    if feedback.corrections:
        doc.add_heading("Corrections", level=1)
        _table(doc, ["What you said", "Better", "Why"],
               [(c.original, c.corrected, c.explanation) for c in feedback.corrections])

    if feedback.improvement_areas:
        doc.add_heading("What to work on", level=1)
        _bullets(doc, feedback.improvement_areas)

    doc.add_heading("Practice before our next lesson", level=1)
    doc.add_paragraph(feedback.practice_task)

    doc.save(str(output_path))


def save_feedback(feedback, folder, lesson_date=None):
    """Write feedback.json (for editing later) and feedback.docx. Returns the .docx path.

    feedback.json is written first, so edits survive even if Word has the .docx open.
    """
    folder = Path(folder)
    (folder / "feedback.json").write_text(feedback.model_dump_json(indent=2), encoding="utf-8")
    docx_path = folder / "feedback.docx"
    create_feedback_doc(feedback, docx_path, lesson_date)
    return docx_path


def load_feedback(folder):
    """The feedback saved in this lesson folder, or None if there isn't any."""
    try:
        return StudentFeedback.model_validate_json((Path(folder) / "feedback.json").read_text(encoding="utf-8"))
    except FileNotFoundError:
        return None


def feedback_as_text(feedback, lesson_date=None):
    """The feedback as plain text, ready to paste into an email or chat message."""
    parts = [f"Lesson feedback – {_date_line(lesson_date)}", "", "What we covered", feedback.lesson_summary]

    def section(title, lines):
        if lines:
            parts.extend(["", title, *lines])

    section("What you did well", [f"• {p}" for p in feedback.positive_feedback])
    section("Grammar", [f"• {g.name}: {g.form}. {g.usage}" for g in feedback.grammar_points])
    section("Vocabulary", [f"• {v.word} – {v.definition} (\"{v.in_context}\")" for v in feedback.vocab_items])
    section("Corrections", [f"• You said: \"{c.original}\" → Better: \"{c.corrected}\". {c.explanation}"
                            for c in feedback.corrections])
    section("What to work on", [f"• {i}" for i in feedback.improvement_areas])
    section("Practice before our next lesson", [feedback.practice_task])
    return "\n".join(parts)
